#!/usr/bin/env python3
"""build_deck.py — generate an investor PowerPoint (.pptx) from a deal-profile JSON.

Builds a partner/investor deck: Title, Core Asset Arbitrage, Value-Creation Playbook (table),
Tech Stack, Competitor Landscape + "different not better" wedge, Financials / Rule of 40, Exit
Plan, and The Ask. The Operating Partner assembles the JSON from the engagement, then runs this.

Requires python-pptx for the actual .pptx:  pip install python-pptx
Use --dry-run to preview the slide plan WITHOUT python-pptx installed.

Examples:
  python3 build_deck.py --content deck-content.json --out deal.pptx
  python3 build_deck.py --content deck-content.json --dry-run
"""
import argparse
import json
import sys


def load_content(path):
    with open(path) as f:
        return json.load(f)


def slide_plan(c):
    """Return an ordered list of (title, list-of-lines) describing each slide."""
    plan = []
    t = c.get("title", {})
    plan.append((t.get("title", "Acquisition Brief"),
                 [t.get("subtitle", ""), t.get("date", "")]))

    caa = c.get("core_asset_arbitrage")
    if caa:
        st = caa.get("current_state", {})
        plan.append(("Core Asset Arbitrage", [
            f"Purchase price: {caa.get('purchase_price', 'TBD')}",
            f"Current state: {st.get('mrr','?')} MRR | {st.get('paying_users','?')} paying users | "
            f"{st.get('hosting_cost','?')} hosting",
            f"Underlying value: {caa.get('underlying_value','Code built + validated paying users')}",
        ]))

    vc = c.get("value_creation")
    if vc:
        lines = ["Lever | Current Flaw | Tactical Move | Projected Result"]
        for row in vc:
            lines.append(f"{row.get('lever','')} | {row.get('flaw','')} | "
                         f"{row.get('move','')} | {row.get('result','')}")
        plan.append(("Value-Creation Playbook", lines))

    ts = c.get("tech_stack")
    if ts:
        plan.append(("Tech Stack", ts if isinstance(ts, list) else [str(ts)]))

    comp = c.get("competitors")
    if comp:
        lines = ["Competitor | Pricing | ICP | Gap"]
        for row in comp.get("matrix", []):
            lines.append(f"{row.get('name','')} | {row.get('pricing','')} | "
                         f"{row.get('icp','')} | {row.get('gap','')}")
        lines.append("")
        lines.append(f"OUR WEDGE (different, not better): {comp.get('wedge','TBD')}")
        plan.append(("Competitor Landscape & Positioning", lines))

    fin = c.get("financials")
    if fin:
        plan.append(("Financials & Rule of 40",
                     [f"{k}: {v}" for k, v in fin.items()]))

    ex = c.get("exit")
    if ex:
        plan.append(("Exit Plan",
                     [f"{k}: {v}" for k, v in ex.items()]))

    ask = c.get("ask")
    if ask:
        plan.append(("The Ask", ask if isinstance(ask, list) else [str(ask)]))

    return plan


def dry_run(plan):
    print("\nSlide plan ({} slides):".format(len(plan)))
    print("=" * 50)
    for i, (title, lines) in enumerate(plan, 1):
        body = [ln for ln in lines if ln]
        print(f"{i:>2}. {title}  ({len(body)} content line(s))")
        for ln in body:
            print(f"      • {ln}")
    print()


def build_pptx(plan, out):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        sys.stderr.write(
            "\npython-pptx is not installed — cannot build the .pptx.\n"
            "Install it once:  pip install python-pptx\n"
            "Or preview the deck now with:  --dry-run\n\n")
        return 2

    prs = Presentation()
    blank = prs.slide_layouts[6]
    title_only = prs.slide_layouts[5]

    def is_table(lines):
        return lines and "|" in lines[0]

    for idx, (title, lines) in enumerate(plan):
        body = [ln for ln in lines if ln]
        if idx == 0:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title
            if slide.placeholders and len(slide.placeholders) > 1:
                slide.placeholders[1].text = "  ".join(body)
            continue

        slide = prs.slides.add_slide(title_only)
        slide.shapes.title.text = title

        if is_table(lines):
            rows_data = [ln.split("|") for ln in body]
            n_rows = len(rows_data)
            n_cols = max(len(r) for r in rows_data)
            tbl = slide.shapes.add_table(
                n_rows, n_cols, Inches(0.4), Inches(1.6), Inches(9.2), Inches(0.4 * n_rows)
            ).table
            for ri, row in enumerate(rows_data):
                for ci in range(n_cols):
                    cell = tbl.cell(ri, ci)
                    cell.text = row[ci].strip() if ci < len(row) else ""
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(11)
                            if ri == 0:
                                run.font.bold = True
        else:
            box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(9), Inches(5))
            tf = box.text_frame
            tf.word_wrap = True
            for j, ln in enumerate(body):
                para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                para.text = "• " + ln
                for run in para.runs:
                    run.font.size = Pt(16)

    prs.save(out)
    print(f"✅ Wrote {out}  ({len(plan)} slides)")
    return 0


def build_parser():
    p = argparse.ArgumentParser(
        description="Generate an investor .pptx from a deal-profile JSON.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("--content", required=True, help="Path to deal-profile JSON")
    p.add_argument("--out", default="deal-deck.pptx", help="Output .pptx path")
    p.add_argument("--dry-run", action="store_true",
                   help="Print the slide plan without building (no python-pptx needed)")
    return p


def main(argv=None):
    args = build_parser().parse_args(argv)
    plan = slide_plan(load_content(args.content))
    if args.dry_run:
        dry_run(plan)
        return 0
    return build_pptx(plan, args.out)


if __name__ == "__main__":
    sys.exit(main())
